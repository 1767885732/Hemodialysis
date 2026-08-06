# -*- coding: utf-8 -*-
"""将账号数据导出为 Excel（.xlsx），包含从者与礼装两个工作表。"""
from __future__ import annotations

from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# 表头定义：(字段键, 表头名, 列宽)
SERVANT_COLUMNS = [
    ("id", "从者ID", 10),
    ("name", "名称", 26),
    ("class", "职阶", 12),
    ("level", "等级", 8),
    ("skill1", "技能1", 8),
    ("skill2", "技能2", 8),
    ("skill3", "技能3", 8),
    ("np_level", "宝具等级", 10),
    ("bond", "羁绊等级", 10),
    ("count", "持有数", 8),
]

EQUIP_COLUMNS = [
    ("id", "礼装ID", 10),
    ("name", "名称", 30),
    ("level", "等级", 8),
    ("count", "持有数", 10),
]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CELL_ALIGN = Alignment(vertical="center", horizontal="center")
TITLE_FONT = Font(bold=True, size=14, color="1F3864")


def _write_sheet(ws, title: str, columns, rows: List[Dict[str, Any]]) -> None:
    # 标题行
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=len(columns)
    )
    ws.cell(row=1, column=1).alignment = Alignment(
        vertical="center", horizontal="left"
    )
    ws.row_dimensions[1].height = 22

    # 表头行
    header_row = 2
    for col_idx, (_key, header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CELL_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[header_row].height = 20

    # 数据行
    for r_off, row_data in enumerate(rows, start=1):
        for col_idx, (key, _header, _width) in enumerate(columns, start=1):
            value = row_data.get(key)
            if value is None:
                value = ""
            cell = ws.cell(row=header_row + r_off, column=col_idx, value=value)
            cell.alignment = CELL_ALIGN

    # 冻结表头
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    # 自动筛选
    if rows:
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}"
        )


def export_to_excel(
    servants: List[Dict[str, Any]],
    equips: List[Dict[str, Any]],
    file_path: str,
    account_label: Optional[str] = None,
) -> str:
    """把从者与礼装列表写入 Excel 文件，返回文件路径。

    参数:
        servants: 从者列表（dict 字段见 SERVANT_COLUMNS）
        equips:   礼装列表（dict 字段见 EQUIP_COLUMNS）
        file_path: 输出 .xlsx 路径
        account_label: 可选账号标识，写入从者表标题
    """
    wb = Workbook()

    ws_svt = wb.active
    ws_svt.title = "从者"
    label = f"FGO 账号从者一览（共 {len(servants)} 位）"
    if account_label:
        label += f" - {account_label}"
    _write_sheet(ws_svt, label, SERVANT_COLUMNS, servants)

    ws_eq = wb.create_sheet(title="礼装")
    _write_sheet(
        ws_eq,
        f"FGO 账号礼装一览（共 {len(equips)} 种）",
        EQUIP_COLUMNS,
        equips,
    )

    wb.save(file_path)
    return file_path
