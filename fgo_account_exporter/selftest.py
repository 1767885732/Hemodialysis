# -*- coding: utf-8 -*-
"""快速自测：验证导出与解析逻辑（不依赖网络与真实凭证）。"""
from __future__ import annotations

import json
import os

from excel_export import export_to_excel
from fgo_api import FGOClient
from sample_data import sample_equips, sample_servants


def test_export_sample():
    out = "/workspace/fgo_account_exporter/test_sample.xlsx"
    path = export_to_excel(sample_servants(), sample_equips(), out, account_label="CN")
    assert os.path.exists(path)
    size = os.path.getsize(path)
    print(f"[OK] 导出样本 Excel 成功: {path} ({size} bytes)")
    # 校验内容
    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert "从者" in wb.sheetnames and "礼装" in wb.sheetnames
    ws = wb["从者"]
    # 标题在第1行，表头第2行，数据从第3行开始
    servant_rows = ws.max_row - 2
    eq_rows = wb["礼装"].max_row - 2
    assert servant_rows == len(sample_servants()), servant_rows
    assert eq_rows == len(sample_equips()), eq_rows
    print(f"[OK] 工作表校验通过：从者 {servant_rows} 行，礼装 {eq_rows} 行")


def test_parse_mock():
    """构造一个贴近真实 /home/top 的响应，验证解析。"""
    mock = {
        "response": [
            {
                "resCode": "00",
                "resMsg": "",
                "cache": {
                    "updated": {
                        "mstSvt": [
                            {"id": 100100, "name": "阿尔托莉雅·潘德拉贡"},
                            {"id": 701300, "name": "赫拉克勒斯"},
                        ],
                        "mstEquip": [
                            {"id": 9400190, "name": "万华镜"},
                            {"id": 9400340, "name": "黑之圣杯"},
                        ],
                        "userServant": [
                            {"svtId": 100100, "classId": 1,
                             "status": {"lv": 90, "skillLv1": 10, "skillLv2": 10,
                                        "skillLv3": 10, "treasureDeviceLv": 5,
                                        "friendshipRank": 15}},
                            {"svtId": 701300, "classId": 7,
                             "status": {"lv": 80, "skillLv1": 10, "skillLv2": 9,
                                        "skillLv3": 9, "treasureDeviceLv": 4,
                                        "friendshipRank": 11}},
                        ],
                        "userEquip": [
                            {"svtId": 9400190, "status": {"lv": 100}},
                            {"svtId": 9400190, "status": {"lv": 95}},
                            {"svtId": 9400340, "status": {"lv": 100}},
                        ],
                        "userEquipCollection": [
                            {"svtId": 9400190, "num": 2},
                            {"svtId": 9400340, "num": 1},
                        ],
                    }
                },
            }
        ]
    }

    client = FGOClient(region="CN", user_id="1", auth_key="a", secret_key="b")
    svt_names, eq_names = client.build_name_maps(mock)
    assert svt_names[100100] == "阿尔托莉雅·潘德拉贡"
    assert eq_names[9400190] == "万华镜"

    servants = client.parse_servants(mock)
    assert len(servants) == 2, servants
    s = next(s for s in servants if s["id"] == 100100)
    assert s["name"] == "阿尔托莉雅·潘德拉贡"
    assert s["level"] == 90 and s["np_level"] == 5 and s["class"] == "Saber"
    h = next(s for s in servants if s["id"] == 701300)
    assert h["class"] == "Berserker"

    equips = client.parse_equips(mock)
    by_id = {e["id"]: e for e in equips}
    assert by_id[9400190]["name"] == "万华镜"
    assert by_id[9400190]["count"] == 2, by_id[9400190]
    assert by_id[9400340]["count"] == 1
    print(f"[OK] 解析自测通过：从者 {len(servants)} 位，礼装 {len(equips)} 种")

    # 导出解析结果
    out = "/workspace/fgo_account_exporter/test_parsed.xlsx"
    export_to_excel(servants, equips, out, account_label="CN(模拟)")
    assert os.path.exists(out)
    print(f"[OK] 模拟解析结果导出成功: {out}")


if __name__ == "__main__":
    test_export_sample()
    test_parse_mock()
    print("\n全部自测通过 ✅")
